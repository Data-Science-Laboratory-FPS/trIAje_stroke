"""Descriptive and inferential paired comparison of model and telephone triage.

All comparisons in this module are restricted to the same calls. The module
includes paired tests, bootstrap inference, sex-gap analyses, and matched
operating-point sensitivity analyses.
"""

import os

import numpy as np
import pandas as pd
from scipy.stats import beta, binomtest, chi2, norm
from sklearn.metrics import confusion_matrix, roc_curve


_REQUIRED_METRICS = (
    "TN",
    "FP",
    "FN",
    "TP",
    "Recall",
    "Specificity",
    "PPV",
    "NPV",
    "Overtriage",
    "Undertriage",
    "P1 predicted",
)

_PROPORTION_METRICS = (
    "Recall",
    "Specificity",
    "PPV",
    "NPV",
    "Overtriage",
    "Undertriage",
)

_DISPLAY_METRIC_LABELS = {"Recall": "Sensitivity"}


def get_common_comparison_subset(
    df,
    y_col="target_real",
    model_col="target_pred_model",
    triage_col="p1_assigned",
):
    """Return rows with complete binary outcomes and predictions.

    The input index is preserved.  A report from an earlier call is retained
    when an already-created common subset is passed downstream; this lets every
    comparative output use the same exclusion report without filtering again.
    """
    required = [y_col, model_col, triage_col]
    missing_columns = [column for column in required if column not in df.columns]
    if missing_columns:
        raise KeyError(
            "Common comparison requires columns: "
            + ", ".join(required)
            + ". Missing: "
            + ", ".join(missing_columns)
        )

    complete_mask = df[required].notna().all(axis=1)
    common_df = df.loc[complete_mask].copy()
    excluded_df = df.loc[~complete_mask]

    for column in required:
        numeric_values = pd.to_numeric(common_df[column], errors="raise")
        assert set(numeric_values.unique()).issubset({0, 1}), (
            "Column {!r} contains values outside {{0, 1}}: {}".format(
                column, sorted(numeric_values.unique().tolist())
            )
        )
        common_df[column] = numeric_values.astype("int8")

    excluded_outcome_distribution = (
        df.loc[~complete_mask, y_col]
        .value_counts(dropna=False)
        .to_dict()
    )
    report = {
        "n_input": int(len(df)),
        "n_common": int(complete_mask.sum()),
        "n_excluded": int((~complete_mask).sum()),
        "excluded_indices": excluded_df.index.tolist(),
        "excluded_outcome_distribution": excluded_outcome_distribution,
    }

    # Integration deliberately passes the common frame through several
    # consumers.  Preserve its original report so captions still state how many
    # test calls were excluded, even if this helper is called again downstream.
    inherited_report = df.attrs.get("common_subset_report")
    if report["n_excluded"] == 0 and inherited_report is not None:
        report = dict(inherited_report)
    common_df.attrs["common_subset_report"] = report
    return common_df, report


def _clopper_pearson(successes, trials, alpha=0.05):
    """Clopper-Pearson exact binomial interval, as proportions."""
    if trials == 0:
        return np.nan, np.nan
    low = (
        float(beta.ppf(alpha / 2, successes, trials - successes + 1))
        if successes > 0 else 0.0
    )
    high = (
        float(beta.ppf(1 - alpha / 2, successes + 1, trials - successes))
        if successes < trials else 1.0
    )
    return low, high


def _proportion(successes, trials):
    if trials == 0:
        return np.nan, (np.nan, np.nan)
    value = successes / trials
    return value, _clopper_pearson(int(successes), int(trials))


def _wilson_interval(successes, trials, alpha=0.05):
    """Wilson score interval for one binomial proportion."""
    if trials == 0:
        return np.nan, np.nan
    z_value = norm.ppf(1 - alpha / 2)
    proportion = successes / trials
    denominator = 1 + z_value ** 2 / trials
    centre = (proportion + z_value ** 2 / (2 * trials)) / denominator
    half_width = (
        z_value
        * np.sqrt(
            proportion * (1 - proportion) / trials
            + z_value ** 2 / (4 * trials ** 2)
        )
        / denominator
    )
    return max(-1.0, centre - half_width), min(1.0, centre + half_width)


def newcombe_paired_interval(a, b, c, d, alpha=0.05):
    """Newcombe (1998) method 10 interval for a paired difference."""
    n = a + b + c + d
    if n == 0:
        return np.nan, np.nan, np.nan, np.nan
    p_model = (a + b) / n
    p_triage = (a + c) / n
    difference = p_model - p_triage
    model_low, model_high = _wilson_interval(a + b, n, alpha=alpha)
    triage_low, triage_high = _wilson_interval(a + c, n, alpha=alpha)
    phi_denominator = np.sqrt(
        (a + b) * (c + d) * (a + c) * (b + d)
    )
    phi = (a * d - b * c) / phi_denominator if phi_denominator else 0.0
    lower_term = (
        (p_model - model_low) ** 2
        - 2 * phi * (p_model - model_low) * (triage_high - p_triage)
        + (triage_high - p_triage) ** 2
    )
    upper_term = (
        (model_high - p_model) ** 2
        - 2 * phi * (model_high - p_model) * (p_triage - triage_low)
        + (p_triage - triage_low) ** 2
    )
    lower = difference - np.sqrt(max(0.0, lower_term))
    upper = difference + np.sqrt(max(0.0, upper_term))
    return max(-1.0, lower), min(1.0, upper), difference, phi


def newcombe_independent_interval(k1, n1, k2, n2, alpha=0.05):
    """Newcombe interval for the difference of independent proportions."""
    if n1 == 0 or n2 == 0:
        return np.nan, np.nan
    p1 = k1 / n1
    p2 = k2 / n2
    l1, u1 = _wilson_interval(k1, n1, alpha=alpha)
    l2, u2 = _wilson_interval(k2, n2, alpha=alpha)
    difference = p1 - p2
    lower = difference - np.sqrt((p1 - l1) ** 2 + (u2 - p2) ** 2)
    upper = difference + np.sqrt((u1 - p1) ** 2 + (p2 - l2) ** 2)
    return max(-1.0, lower), min(1.0, upper)


def compare_paired_proportion(correct_model, correct_triage, label):
    """Compare paired correctness indicators with McNemar and Newcombe."""
    model = np.asarray(correct_model, dtype=np.int8)
    triage = np.asarray(correct_triage, dtype=np.int8)
    if model.shape != triage.shape:
        raise ValueError("Paired indicators must have the same shape.")
    assert set(np.unique(model)).issubset({0, 1})
    assert set(np.unique(triage)).issubset({0, 1})

    a = int(np.sum((model == 1) & (triage == 1)))
    b = int(np.sum((model == 1) & (triage == 0)))
    c = int(np.sum((model == 0) & (triage == 1)))
    d = int(np.sum((model == 0) & (triage == 0)))
    discordant = b + c
    exact_p = float(binomtest(b, discordant, 0.5).pvalue) if discordant else 1.0
    if discordant:
        corrected_statistic = max(0, abs(b - c) - 1) ** 2 / discordant
        uncorrected_statistic = (b - c) ** 2 / discordant
        corrected_p = float(chi2.sf(corrected_statistic, 1))
        uncorrected_p = float(chi2.sf(uncorrected_statistic, 1))
    else:
        corrected_statistic = uncorrected_statistic = 0.0
        corrected_p = uncorrected_p = 1.0
    test_reported = "Exact binomial McNemar" if discordant < 25 else "Continuity-corrected McNemar chi-square"
    reported_p = exact_p if discordant < 25 else corrected_p
    ci_low, ci_high, difference, phi = newcombe_paired_interval(a, b, c, d)
    return {
        "label": label,
        "a": a,
        "b": b,
        "c": c,
        "d": d,
        "n": int(len(model)),
        "difference": difference,
        "ci": (ci_low, ci_high),
        "phi": phi,
        "discordant": discordant,
        "exact_p_value": exact_p,
        "chi2_corrected_statistic": corrected_statistic,
        "chi2_corrected_p_value": corrected_p,
        "chi2_uncorrected_statistic": uncorrected_statistic,
        "chi2_uncorrected_p_value": uncorrected_p,
        "test_reported": test_reported,
        "p_value": reported_p,
    }


def generate_outcome_year_bootstrap_indices(
    y, year, n_bootstrap=1000, seed=42
):
    """Generate paired positional resamples preserving every outcome-year cell."""
    from kbase.modeling import _make_outcome_year_strata

    y_series = pd.Series(np.asarray(y)).reset_index(drop=True)
    year_series = pd.Series(np.asarray(year)).reset_index(drop=True)
    strata = _make_outcome_year_strata(y_series, year_series)
    stratum_positions = [
        group.index.to_numpy(dtype=int)
        for _, group in strata.groupby(strata, sort=False)
    ]
    rng = np.random.default_rng(seed)
    return [
        np.concatenate([
            rng.choice(positions, size=len(positions), replace=True)
            for positions in stratum_positions
        ])
        for _ in range(n_bootstrap)
    ]


def _bootstrap_interval(values):
    values = np.asarray(values, dtype=float)
    valid = values[~np.isnan(values)]
    if len(valid) < len(values) / 2:
        print(
            "WARNING: only {}/{} valid paired bootstrap replicates; "
            "the interval may be unreliable.".format(len(valid), len(values))
        )
    if not len(valid):
        return (np.nan, np.nan), 0
    return (
        float(np.percentile(valid, 2.5)),
        float(np.percentile(valid, 97.5)),
    ), int(len(valid))


def _bootstrap_two_sided_p(values):
    valid = np.asarray(values, dtype=float)
    valid = valid[~np.isnan(valid)]
    if not len(valid):
        return np.nan
    probability_nonpositive = np.mean(valid <= 0)
    probability_nonnegative = np.mean(valid >= 0)
    p_value = min(1.0, 2 * min(probability_nonpositive, probability_nonnegative))
    return max(1 / (len(valid) + 1), p_value)


def _prediction_error_rate(y, prediction, error_type):
    y = np.asarray(y)
    prediction = np.asarray(prediction)
    if error_type == "undertriage":
        selected = prediction == 0
        errors = y[selected] == 1
    elif error_type == "overtriage":
        selected = prediction == 1
        errors = y[selected] == 0
    else:
        raise ValueError("error_type must be 'undertriage' or 'overtriage'.")
    denominator = int(selected.sum())
    if denominator == 0:
        return np.nan, 0
    return float(errors.mean()), denominator


def _system_metrics(y_true, y_pred):
    tn, fp, fn, tp = confusion_matrix(y_true, y_pred, labels=[0, 1]).ravel()
    tn, fp, fn, tp = (int(tn), int(fp), int(fn), int(tp))

    recall, recall_ci = _proportion(tp, tp + fn)
    specificity, specificity_ci = _proportion(tn, tn + fp)
    ppv, ppv_ci = _proportion(tp, tp + fp)
    npv, npv_ci = _proportion(tn, tn + fn)
    overtriage, overtriage_ci = _proportion(fp, tp + fp)
    undertriage, undertriage_ci = _proportion(fn, tn + fn)

    return {
        "TN": tn,
        "FP": fp,
        "FN": fn,
        "TP": tp,
        "Recall": recall,
        "Recall CI": recall_ci,
        "Specificity": specificity,
        "Specificity CI": specificity_ci,
        "PPV": ppv,
        "PPV CI": ppv_ci,
        "NPV": npv,
        "NPV CI": npv_ci,
        "Overtriage": overtriage,
        "Overtriage CI": overtriage_ci,
        "Undertriage": undertriage,
        "Undertriage CI": undertriage_ci,
        "P1 predicted": tp + fp,
    }


def _statsmodels_available():
    try:
        import statsmodels.api  # noqa: F401
        return True
    except ImportError:
        return False


def _gee_paired_error_difference(y, model_prediction, triage_prediction, error_type):
    """Fit the marginal paired GEE comparison requested for prediction errors."""
    import statsmodels.api as sm

    frames = []
    for system_indicator, prediction in (
        (0, np.asarray(triage_prediction)),
        (1, np.asarray(model_prediction)),
    ):
        y_array = np.asarray(y)
        selected = prediction == (0 if error_type == "undertriage" else 1)
        response = y_array if error_type == "undertriage" else 1 - y_array
        positions = np.flatnonzero(selected)
        frames.append(pd.DataFrame({
            "patient_id": positions,
            "response": response[selected].astype(float),
            "model_system": np.full(len(positions), system_indicator, dtype=float),
        }))
    long_df = pd.concat(frames, ignore_index=True)
    design = sm.add_constant(long_df[["model_system"]], has_constant="add")
    try:
        identity_family = sm.families.Binomial(
            link=sm.families.links.Identity()
        )
        fit = sm.GEE(
            long_df["response"],
            design,
            groups=long_df["patient_id"],
            family=identity_family,
            cov_struct=sm.cov_struct.Independence(),
        ).fit()
        if not getattr(fit, "converged", True):
            raise RuntimeError("Identity-link GEE did not converge.")
        estimate = float(fit.params["model_system"])
        standard_error = float(fit.bse["model_system"])
        z_value = estimate / standard_error if standard_error else 0.0
        return {
            "method": "GEE Wald z-test (identity link)",
            "difference": estimate,
            "standard_error": standard_error,
            "z_value": z_value,
            "p_value": float(2 * norm.sf(abs(z_value))),
            "odds_ratio": np.nan,
            "odds_ratio_ci": (np.nan, np.nan),
        }
    except Exception as error:
        print(
            "WARNING: identity-link GEE failed for {} ({}); using a logit-link "
            "GEE and reporting the odds ratio.".format(error_type, error)
        )
        fit = sm.GEE(
            long_df["response"],
            design,
            groups=long_df["patient_id"],
            family=sm.families.Binomial(),
            cov_struct=sm.cov_struct.Independence(),
        ).fit()
        coefficient = float(fit.params["model_system"])
        standard_error = float(fit.bse["model_system"])
        z_value = coefficient / standard_error if standard_error else 0.0
        return {
            "method": "GEE Wald z-test (logit fallback)",
            "difference": np.nan,
            "standard_error": standard_error,
            "z_value": z_value,
            "p_value": float(2 * norm.sf(abs(z_value))),
            "odds_ratio": float(np.exp(coefficient)),
            "odds_ratio_ci": (
                float(np.exp(coefficient - 1.96 * standard_error)),
                float(np.exp(coefficient + 1.96 * standard_error)),
            ),
        }


def compare_paired_error_rate(
    y,
    model_prediction,
    triage_prediction,
    error_type,
    bootstrap_indices,
):
    """Compare undertriage or overtriage with a paired bootstrap and GEE."""
    y_array = np.asarray(y)
    model_array = np.asarray(model_prediction)
    triage_array = np.asarray(triage_prediction)
    model_rate, model_denominator = _prediction_error_rate(
        y_array, model_array, error_type
    )
    triage_rate, triage_denominator = _prediction_error_rate(
        y_array, triage_array, error_type
    )
    difference = model_rate - triage_rate
    bootstrap_differences = []
    for positions in bootstrap_indices:
        model_rep, _ = _prediction_error_rate(
            y_array[positions], model_array[positions], error_type
        )
        triage_rep, _ = _prediction_error_rate(
            y_array[positions], triage_array[positions], error_type
        )
        bootstrap_differences.append(model_rep - triage_rep)
    interval, valid_replicates = _bootstrap_interval(bootstrap_differences)

    model_selected = model_array == (0 if error_type == "undertriage" else 1)
    triage_selected = triage_array == (0 if error_type == "undertriage" else 1)
    both_denominators = int(np.sum(model_selected & triage_selected))
    classification_label = (
        "negative" if error_type == "undertriage" else "positive"
    )
    model_classification_proportion = np.mean(
        model_array == (0 if error_type == "undertriage" else 1)
    )
    triage_classification_proportion = np.mean(
        triage_array == (0 if error_type == "undertriage" else 1)
    )
    print(
        "{} denominators: model={}, triage={}, present in both={}; "
        "{} classification proportions: model={:.2%}, triage={:.2%}.".format(
            error_type.capitalize(),
            model_denominator,
            triage_denominator,
            both_denominators,
            classification_label,
            model_classification_proportion,
            triage_classification_proportion,
        )
    )

    if _statsmodels_available():
        test = _gee_paired_error_difference(
            y_array, model_array, triage_array, error_type
        )
        test_method = test["method"]
        p_value = test["p_value"]
    else:
        test = None
        test_method = "Outcome-year stratified paired bootstrap"
        p_value = _bootstrap_two_sided_p(bootstrap_differences)
        print(
            "statsmodels is unavailable; {} uses the paired-bootstrap "
            "two-sided p-value.".format(error_type)
        )
    return {
        "error_type": error_type,
        "model": model_rate,
        "triage": triage_rate,
        "difference": difference,
        "ci": interval,
        "bootstrap_differences": np.asarray(bootstrap_differences),
        "valid_bootstrap_replicates": valid_replicates,
        "model_denominator": model_denominator,
        "triage_denominator": triage_denominator,
        "patients_in_both_denominators": both_denominators,
        "model_negative_proportion": float(np.mean(model_array == 0)),
        "triage_negative_proportion": float(np.mean(triage_array == 0)),
        "test_reported": test_method,
        "p_value": p_value,
        "gee": test,
    }


def calculate_categorical_nri(
    y, model_prediction, triage_prediction, bootstrap_indices=None
):
    """Calculate categorical NRI and its four movement components."""
    y_array = np.asarray(y)
    model = np.asarray(model_prediction)
    triage = np.asarray(triage_prediction)

    def calculate(y_values, model_values, triage_values):
        event = y_values == 1
        nonevent = y_values == 0
        if not event.any() or not nonevent.any():
            return {
                "event_up": np.nan,
                "event_down": np.nan,
                "nonevent_down": np.nan,
                "nonevent_up": np.nan,
                "nri_events": np.nan,
                "nri_nonevents": np.nan,
                "nri_overall": np.nan,
            }
        event_up = float(np.mean(model_values[event] > triage_values[event]))
        event_down = float(np.mean(model_values[event] < triage_values[event]))
        nonevent_down = float(np.mean(model_values[nonevent] < triage_values[nonevent]))
        nonevent_up = float(np.mean(model_values[nonevent] > triage_values[nonevent]))
        nri_events = event_up - event_down
        nri_nonevents = nonevent_down - nonevent_up
        return {
            "event_up": event_up,
            "event_down": event_down,
            "nonevent_down": nonevent_down,
            "nonevent_up": nonevent_up,
            "nri_events": nri_events,
            "nri_nonevents": nri_nonevents,
            "nri_overall": nri_events + nri_nonevents,
        }

    result = calculate(y_array, model, triage)
    model_metrics = _system_metrics(y_array, model)
    triage_metrics = _system_metrics(y_array, triage)
    delta_sensitivity = model_metrics["Recall"] - triage_metrics["Recall"]
    delta_specificity = model_metrics["Specificity"] - triage_metrics["Specificity"]
    delta_youden = (
        model_metrics["Recall"] + model_metrics["Specificity"] - 1
        - (triage_metrics["Recall"] + triage_metrics["Specificity"] - 1)
    )
    assert abs(result["nri_events"] - delta_sensitivity) <= 1e-12
    assert abs(result["nri_nonevents"] - delta_specificity) <= 1e-12
    assert abs(result["nri_overall"] - delta_youden) <= 1e-12

    bootstrap_values = {key: [] for key in (
        "nri_events", "nri_nonevents", "nri_overall"
    )}
    if bootstrap_indices is not None:
        for positions in bootstrap_indices:
            replicate = calculate(
                y_array[positions], model[positions], triage[positions]
            )
            for key in bootstrap_values:
                bootstrap_values[key].append(replicate[key])
        result["ci"] = {
            key: _bootstrap_interval(values)[0]
            for key, values in bootstrap_values.items()
        }
        result["bootstrap_values"] = {
            key: np.asarray(values) for key, values in bootstrap_values.items()
        }
    else:
        result["ci"] = {key: (np.nan, np.nan) for key in bootstrap_values}
        result["bootstrap_values"] = bootstrap_values
    return result


def _metric_fraction(y, prediction, metric):
    y_array = np.asarray(y)
    prediction_array = np.asarray(prediction)
    if metric == "Sensitivity":
        selected = y_array == 1
        success = prediction_array[selected] == 1
    elif metric == "Specificity":
        selected = y_array == 0
        success = prediction_array[selected] == 0
    elif metric == "Undertriage":
        selected = prediction_array == 0
        success = y_array[selected] == 1
    else:
        raise ValueError("Unsupported metric: {}".format(metric))
    denominator = int(selected.sum())
    if denominator == 0:
        return np.nan, 0, 0
    numerator = int(success.sum())
    return numerator / denominator, numerator, denominator


def classify_disparity_change(triage_gap, model_gap):
    """Describe how the absolute sex gap changes between the two systems."""
    if pd.isna(triage_gap) or pd.isna(model_gap):
        return "—"
    if triage_gap * model_gap < 0:
        return "reverses"
    if abs(model_gap) < abs(triage_gap):
        return "narrows"
    if abs(model_gap) > abs(triage_gap):
        return "widens"
    return "unchanged"


def compare_sex_gaps(
    common_df,
    bootstrap_indices,
    model_full_test_df=None,
):
    """Compare Men-Women metric gaps and their difference between systems."""
    if "sex" not in common_df.columns:
        print("WARNING: sex is unavailable; sex-gap analysis was skipped.")
        return {"rows": [], "table": pd.DataFrame(), "n_missing_sex": len(common_df)}
    sex_numeric = pd.to_numeric(common_df["sex"], errors="coerce").to_numpy()
    valid_sex = np.isin(sex_numeric, [0, 1])
    n_missing = int((~valid_sex).sum())
    print(
        "Sex-gap common subset: Men N={:,}, Women N={:,}, excluded missing/other sex={:,}.".format(
            int(np.sum(sex_numeric == 0)), int(np.sum(sex_numeric == 1)), n_missing
        )
    )
    y = common_df["target_real"].to_numpy()
    predictions = {
        "Telephone triage": common_df["p1_assigned"].to_numpy(),
        "ML model": common_df["target_pred_model"].to_numpy(),
    }
    rows = []
    for metric in ("Sensitivity", "Specificity", "Undertriage"):
        system_results = {}
        for system, prediction in predictions.items():
            men_mask = sex_numeric == 0
            women_mask = sex_numeric == 1
            men_value, men_k, men_n = _metric_fraction(
                y[men_mask], prediction[men_mask], metric
            )
            women_value, women_k, women_n = _metric_fraction(
                y[women_mask], prediction[women_mask], metric
            )
            gap = men_value - women_value
            gap_ci = newcombe_independent_interval(men_k, men_n, women_k, women_n)
            reason = ""
            if np.isnan(gap):
                reason = "Undefined because at least one sex-specific denominator is zero."
            system_results[system] = {
                "men": men_value,
                "men_ci": _clopper_pearson(men_k, men_n),
                "women": women_value,
                "women_ci": _clopper_pearson(women_k, women_n),
                "men_n": men_n,
                "women_n": women_n,
                "gap": gap,
                "gap_ci": gap_ci,
                "reason": reason,
            }

        difference_in_differences = (
            system_results["ML model"]["gap"]
            - system_results["Telephone triage"]["gap"]
        )
        bootstrap_dd = []
        for positions in bootstrap_indices:
            replicate_y = y[positions]
            replicate_sex = sex_numeric[positions]
            replicate_gaps = {}
            for system, prediction in predictions.items():
                replicate_prediction = prediction[positions]
                men = replicate_sex == 0
                women = replicate_sex == 1
                men_value = _metric_fraction(
                    replicate_y[men], replicate_prediction[men], metric
                )[0]
                women_value = _metric_fraction(
                    replicate_y[women], replicate_prediction[women], metric
                )[0]
                replicate_gaps[system] = men_value - women_value
            bootstrap_dd.append(
                replicate_gaps["ML model"]
                - replicate_gaps["Telephone triage"]
            )
        dd_ci, valid_replicates = _bootstrap_interval(bootstrap_dd)
        rows.append({
            "metric": metric,
            "triage": system_results["Telephone triage"],
            "model": system_results["ML model"],
            "difference_in_differences": difference_in_differences,
            "difference_in_differences_ci": dd_ci,
            "p_value": _bootstrap_two_sided_p(bootstrap_dd),
            "valid_bootstrap_replicates": valid_replicates,
            "bootstrap_difference_in_differences": np.asarray(bootstrap_dd),
            "disparity_change": classify_disparity_change(
                system_results["Telephone triage"]["gap"],
                system_results["ML model"]["gap"],
            ),
        })

    model_full_values = {}
    if model_full_test_df is not None and "sex" in model_full_test_df.columns:
        full_sex = pd.to_numeric(model_full_test_df["sex"], errors="coerce").to_numpy()
        full_y = model_full_test_df["target_real"].to_numpy()
        full_prediction = model_full_test_df["target_pred_model"].to_numpy()
        for metric in ("Sensitivity", "Specificity", "Undertriage"):
            model_full_values[metric] = {}
            for sex_value, label in ((0, "Men"), (1, "Women")):
                mask = full_sex == sex_value
                full_value = _metric_fraction(full_y[mask], full_prediction[mask], metric)[0]
                common_value = _metric_fraction(
                    y[sex_numeric == sex_value],
                    predictions["ML model"][sex_numeric == sex_value],
                    metric,
                )[0]
                model_full_values[metric][label] = {
                    "full_test": full_value,
                    "common_subset": common_value,
                    "difference": common_value - full_value,
                }
                print(
                    "ML model {} for {}: full test={:.2%}, common subset={:.2%}, "
                    "common-full={:+.2f} pp.".format(
                        metric, label, full_value, common_value,
                        (common_value - full_value) * 100,
                    )
                )

    display_rows = []
    for row in rows:
        display_rows.append({
            "Metric": row["metric"],
            "Telephone triage Men (95% CI)": _format_proportion_ci(
                row["triage"]["men"], row["triage"]["men_ci"]
            ),
            "Telephone triage Women (95% CI)": _format_proportion_ci(
                row["triage"]["women"], row["triage"]["women_ci"]
            ),
            "Telephone triage Men − Women gap (95% CI)": _format_estimate_ci(
                row["triage"]["gap"], row["triage"]["gap_ci"]
            ),
            "ML model Men (95% CI)": _format_proportion_ci(
                row["model"]["men"], row["model"]["men_ci"]
            ),
            "ML model Women (95% CI)": _format_proportion_ci(
                row["model"]["women"], row["model"]["women_ci"]
            ),
            "ML model Men − Women gap (95% CI)": _format_estimate_ci(
                row["model"]["gap"], row["model"]["gap_ci"]
            ),
            "DD model − triage (95% CI)": _format_estimate_ci(
                row["difference_in_differences"],
                row["difference_in_differences_ci"],
            ),
            "p-value": _format_p_value(row["p_value"]),
            "Disparity change": row["disparity_change"],
        })
    return {
        "rows": rows,
        "table": pd.DataFrame(display_rows),
        "n_men": int(np.sum(sex_numeric == 0)),
        "n_women": int(np.sum(sex_numeric == 1)),
        "n_missing_sex": n_missing,
        "model_full_vs_common": model_full_values,
    }


def _select_matched_threshold(y, probabilities, target, metric):
    """Select an operating-point threshold using training data only."""
    y_array = np.asarray(y)
    probabilities_array = np.asarray(probabilities, dtype=float)
    false_positive_rate, true_positive_rate, thresholds = roc_curve(
        y_array, probabilities_array, drop_intermediate=False
    )
    values = 1 - false_positive_rate if metric == "specificity" else true_positive_rate
    valid_thresholds = thresholds[values >= target]
    if not len(valid_thresholds):
        return np.nan
    return (
        float(np.min(valid_thresholds))
        if metric == "specificity"
        else float(np.max(valid_thresholds))
    )


def matched_operating_point_analysis(
    common_df,
    train_threshold,
    y_train=None,
    y_train_threshold_prob=None,
    train_triage=None,
    year_train=None,
    primary_threshold=None,
):
    """Compare systems at OOF training thresholds matched to triage operation."""
    if train_threshold != "oof":
        print("WARNING: matched operating-point analysis requires OOF probabilities; skipped.")
        return {"skipped": True, "reason": "train_threshold is not 'oof'", "table": pd.DataFrame()}
    if any(value is None for value in (
        y_train, y_train_threshold_prob, train_triage, year_train
    )):
        print("WARNING: matched operating-point analysis lacks aligned training data; skipped.")
        return {"skipped": True, "reason": "missing training inputs", "table": pd.DataFrame()}

    train_index = pd.Index(pd.Series(y_train).index)
    train_df = pd.DataFrame(index=train_index)
    train_df["target_real"] = pd.Series(y_train).reindex(train_index)
    train_probabilities = pd.Series(
        np.asarray(y_train_threshold_prob), index=train_index
    )
    train_df["target_pred_model"] = (
        train_probabilities >= float(primary_threshold)
    ).astype("int8")
    train_df["p1_assigned"] = pd.Series(train_triage).reindex(train_index)
    train_df["year"] = pd.Series(year_train).reindex(train_index)
    train_common, train_report = get_common_comparison_subset(train_df)
    aligned_probabilities = train_probabilities.loc[train_common.index].to_numpy()
    train_y = train_common["target_real"].to_numpy()
    train_triage_prediction = train_common["p1_assigned"].to_numpy()
    triage_train_metrics = _system_metrics(train_y, train_triage_prediction)
    tau_specificity = _select_matched_threshold(
        train_y,
        aligned_probabilities,
        triage_train_metrics["Specificity"],
        "specificity",
    )
    tau_sensitivity = _select_matched_threshold(
        train_y,
        aligned_probabilities,
        triage_train_metrics["Recall"],
        "sensitivity",
    )
    print("Matched operating point thresholds selected on OOF training probabilities only:")
    print("  tau_spec = {:.6f}".format(tau_specificity))
    print("  tau_sens = {:.6f}".format(tau_sensitivity))

    test_y = common_df["target_real"].to_numpy()
    test_probability = common_df["target_prob_model"].to_numpy(dtype=float)
    test_triage = common_df["p1_assigned"].to_numpy()
    rows = []
    for label, threshold in (
        ("Iso-specificity", tau_specificity),
        ("Iso-sensitivity", tau_sensitivity),
    ):
        test_prediction = (test_probability >= threshold).astype(np.int8)
        model_metrics = _system_metrics(test_y, test_prediction)
        triage_metrics = _system_metrics(test_y, test_triage)
        for paired_metric, metrics_key, outcome_value, correct_value in (
            ("Sensitivity", "Recall", 1, 1),
            ("Specificity", "Specificity", 0, 0),
        ):
            population = test_y == outcome_value
            correct_model = (
                test_prediction[population] == correct_value
            ).astype(np.int8)
            correct_triage = (
                test_triage[population] == correct_value
            ).astype(np.int8)
            comparison = compare_paired_proportion(
                correct_model,
                correct_triage,
                "{} at {}".format(paired_metric, label),
            )
            rows.append({
                "analysis": label,
                "threshold": threshold,
                "metric": paired_metric,
                "model_value": model_metrics[metrics_key],
                "triage_value": triage_metrics[metrics_key],
                "paired_comparison": comparison,
            })
    display_df = pd.DataFrame([{
        "Analysis": row["analysis"],
        "OOF training threshold": "{:.6f}".format(row["threshold"]),
        "Metric": row["metric"],
        "Telephone triage": "{:.2f}%".format(row["triage_value"] * 100),
        "ML model": "{:.2f}%".format(row["model_value"] * 100),
        "Δ model − triage (95% CI)": _format_estimate_ci(
            row["paired_comparison"]["difference"],
            row["paired_comparison"]["ci"],
        ),
        "Test": row["paired_comparison"]["test_reported"],
        "p-value": _format_p_value(row["paired_comparison"]["p_value"]),
    } for row in rows])
    return {
        "skipped": False,
        "train_report": train_report,
        "triage_train_metrics": triage_train_metrics,
        "tau_spec": tau_specificity,
        "tau_sens": tau_sensitivity,
        "rows": rows,
        "table": display_df,
    }


def _format_estimate_ci(value, interval):
    if pd.isna(value):
        return "NA"
    low, high = interval
    if pd.isna(low) or pd.isna(high):
        return "{:+.2f} pp".format(value * 100)
    return "{:+.2f} pp ({:+.2f} to {:+.2f})".format(
        value * 100, low * 100, high * 100
    )


def _format_proportion_ci(value, interval):
    if pd.isna(value):
        return "NA"
    low, high = interval
    if pd.isna(low) or pd.isna(high):
        return "{:.2f}%".format(value * 100)
    return "{:.2f}% ({:.2f}–{:.2f})".format(
        value * 100, low * 100, high * 100
    )


def _format_p_value(value):
    if value is None or pd.isna(value):
        return ""
    if value < 0.001:
        return "<0.001"
    return "{:.3f}".format(value)


def _format_system_value(metrics, metric):
    value = metrics[metric]
    if metric not in _PROPORTION_METRICS:
        return "{:,}".format(int(value))
    low, high = metrics[metric + " CI"]
    return "{:.2f}% ({:.2f}–{:.2f})".format(
        value * 100, low * 100, high * 100
    )


def _format_difference(value, metric):
    if metric in _PROPORTION_METRICS:
        return "{:+.2f} pp".format(value * 100)
    return "{:+,d}".format(int(value))


def _build_display_table(
    triage_metrics, model_metrics, differences, inference=None
):
    rows = []
    for metric in _REQUIRED_METRICS:
        difference_text = _format_difference(differences[metric], metric)
        test_text = ""
        p_text = ""
        if inference is not None and metric in ("Recall", "Specificity"):
            result_key = "sensitivity" if metric == "Recall" else "specificity"
            result = inference[result_key]
            difference_text = _format_estimate_ci(result["difference"], result["ci"])
            test_text = result["test_reported"]
            p_text = _format_p_value(result["p_value"])
        elif inference is not None and metric in ("Undertriage", "Overtriage"):
            result = inference[metric.lower()]
            difference_text = _format_estimate_ci(result["difference"], result["ci"])
            test_text = result["test_reported"]
            p_text = _format_p_value(result["p_value"])
        rows.append({
            "Metric": _DISPLAY_METRIC_LABELS.get(metric, metric),
            "Telephone triage (95% CI)": _format_system_value(
                triage_metrics, metric
            ),
            "ML model (95% CI)": _format_system_value(model_metrics, metric),
            "Δ model − triage (95% CI)": difference_text,
            "Test": test_text,
            "p-value": p_text,
        })
    if inference is not None:
        for metric_label, key in (
            ("NRI events", "nri_events"),
            ("NRI non-events", "nri_nonevents"),
            ("NRI overall", "nri_overall"),
        ):
            rows.append({
                "Metric": metric_label,
                "Telephone triage (95% CI)": "",
                "ML model (95% CI)": "",
                "Δ model − triage (95% CI)": _format_estimate_ci(
                    inference["nri"][key], inference["nri"]["ci"][key]
                ),
                "Test": "Categorical NRI",
                "p-value": "",
            })
    return pd.DataFrame(rows)


def _export_word_table(display_df, title, footnote, output_path):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def set_cell_background(cell, colour):
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), colour)
        properties.append(shading)

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)

    title_run = document.add_paragraph().add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(11)
    document.add_paragraph()

    table = document.add_table(rows=len(display_df) + 1, cols=len(display_df.columns))
    table.style = "Table Grid"
    for column_number, column_name in enumerate(display_df.columns):
        cell = table.rows[0].cells[column_number]
        cell.text = str(column_name)
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(8)
        set_cell_background(cell, "2D2D2D")

    for row_number, values in enumerate(display_df.itertuples(index=False)):
        background = "EFEFEF" if row_number % 2 else "FFFFFF"
        for column_number, value in enumerate(values):
            cell = table.rows[row_number + 1].cells[column_number]
            cell.text = str(value)
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            set_cell_background(cell, background)

    document.add_paragraph()
    footnote_run = document.add_paragraph().add_run(footnote)
    footnote_run.italic = True
    footnote_run.font.size = Pt(8)
    document.save(output_path)


def run_paired_inference(
    df,
    n_bootstrap=1000,
    seed=42,
    model_full_test_df=None,
    train_threshold=None,
    y_train=None,
    y_train_threshold_prob=None,
    train_triage=None,
    year_train=None,
    primary_threshold=None,
):
    """Run all paired inferential analyses on the single common subset."""
    common_df, report = get_common_comparison_subset(df)
    if "year" not in common_df.columns:
        raise ValueError("Paired inference requires the test-set year column.")
    y = common_df["target_real"].to_numpy()
    model = common_df["target_pred_model"].to_numpy()
    triage = common_df["p1_assigned"].to_numpy()
    bootstrap_indices = generate_outcome_year_bootstrap_indices(
        y, common_df["year"].to_numpy(), n_bootstrap=n_bootstrap, seed=seed
    )

    event = y == 1
    nonevent = y == 0
    sensitivity = compare_paired_proportion(
        (model[event] == 1).astype(np.int8),
        (triage[event] == 1).astype(np.int8),
        "Sensitivity",
    )
    specificity = compare_paired_proportion(
        (model[nonevent] == 0).astype(np.int8),
        (triage[nonevent] == 0).astype(np.int8),
        "Specificity",
    )
    undertriage = compare_paired_error_rate(
        y, model, triage, "undertriage", bootstrap_indices
    )
    overtriage = compare_paired_error_rate(
        y, model, triage, "overtriage", bootstrap_indices
    )
    nri = calculate_categorical_nri(y, model, triage, bootstrap_indices)
    sex_gaps = compare_sex_gaps(
        common_df, bootstrap_indices, model_full_test_df=model_full_test_df
    )
    matched = matched_operating_point_analysis(
        common_df,
        train_threshold=train_threshold,
        y_train=y_train,
        y_train_threshold_prob=y_train_threshold_prob,
        train_triage=train_triage,
        year_train=year_train,
        primary_threshold=primary_threshold,
    )

    print("\n--- Paired inference summary (ML model − Telephone triage) ---")
    for label, result in (
        ("Sensitivity", sensitivity),
        ("Specificity", specificity),
        ("Undertriage", undertriage),
        ("Overtriage", overtriage),
    ):
        print(
            "{}: delta={:+.2f} pp, 95% CI [{:+.2f}, {:+.2f}], "
            "{} p={}.".format(
                label,
                result["difference"] * 100,
                result["ci"][0] * 100,
                result["ci"][1] * 100,
                result["test_reported"],
                _format_p_value(result["p_value"]),
            )
        )
    print(
        "NRI components: event up={:.2%}, event down={:.2%}, "
        "non-event down={:.2%}, non-event up={:.2%}.".format(
            nri["event_up"], nri["event_down"],
            nri["nonevent_down"], nri["nonevent_up"],
        )
    )
    print(
        "NRI events={:+.2f} pp, NRI non-events={:+.2f} pp, "
        "NRI overall={:+.2f} pp.".format(
            nri["nri_events"] * 100,
            nri["nri_nonevents"] * 100,
            nri["nri_overall"] * 100,
        )
    )
    return {
        "report": report,
        "bootstrap_indices": bootstrap_indices,
        "n_bootstrap": n_bootstrap,
        "seed": seed,
        "sensitivity": sensitivity,
        "specificity": specificity,
        "undertriage": undertriage,
        "overtriage": overtriage,
        "nri": nri,
        "sex_gaps": sex_gaps,
        "matched_operating_point": matched,
    }


def compare_model_vs_triage_common(
    df,
    figures_dir=None,
    export_tables=True,
    run_inference=True,
    n_bootstrap=1000,
    seed=42,
    model_full_test_df=None,
    train_threshold=None,
    y_train=None,
    y_train_threshold_prob=None,
    train_triage=None,
    year_train=None,
    primary_threshold=None,
):
    """Describe model and triage performance on their common test subset."""
    common_df, report = get_common_comparison_subset(df)
    print(
        "\n--- Model vs telephone triage: common test subset ---\n"
        "Common N: {:,}\nExcluded: {:,}".format(
            report["n_common"], report["n_excluded"]
        )
    )

    y_true = common_df["target_real"].to_numpy()
    triage_metrics = _system_metrics(y_true, common_df["p1_assigned"].to_numpy())
    model_metrics = _system_metrics(y_true, common_df["target_pred_model"].to_numpy())
    differences = {
        metric: model_metrics[metric] - triage_metrics[metric]
        for metric in _REQUIRED_METRICS
    }
    operational_totals = {
        "model_p1_predicted": model_metrics["P1 predicted"],
        "triage_p1_assigned": triage_metrics["P1 predicted"],
        "model_fp": model_metrics["FP"],
        "triage_fp": triage_metrics["FP"],
        "model_tp": model_metrics["TP"],
        "triage_tp": triage_metrics["TP"],
        "additional_model_fp": model_metrics["FP"] - triage_metrics["FP"],
        "additional_model_tp": model_metrics["TP"] - triage_metrics["TP"],
    }

    paired_inference = None
    if run_inference:
        paired_inference = run_paired_inference(
            common_df,
            n_bootstrap=n_bootstrap,
            seed=seed,
            model_full_test_df=model_full_test_df,
            train_threshold=train_threshold,
            y_train=y_train,
            y_train_threshold_prob=y_train_threshold_prob,
            train_triage=train_triage,
            year_train=year_train,
            primary_threshold=primary_threshold,
        )
    display_df = _build_display_table(
        triage_metrics, model_metrics, differences, inference=paired_inference
    )
    title = (
        "Comparison of the model and the telephone triage system "
        "(common test subset, N = {:,})".format(report["n_common"])
    )
    if paired_inference is None:
        footnote = (
            "Both systems evaluated on the same {:,} calls of the test set with an "
            "assigned triage priority ({:,} calls excluded from this comparison "
            "only). Individual metrics: Clopper–Pearson 95% CI."
        ).format(report["n_common"], report["n_excluded"])
    else:
        error_test = (
            "GEE Wald z-test with cluster-robust variance"
            if _statsmodels_available()
            else "paired-bootstrap two-sided test"
        )
        footnote = (
            "Both systems evaluated on the same {:,} calls of the test set with an "
            "assigned triage priority ({:,} calls excluded from this comparison "
            "only). Individual metrics: Clopper–Pearson 95% CI. Sensitivity and "
            "specificity differences: Newcombe method 10 paired 95% CI and McNemar "
            "test. Undertriage and overtriage differences: outcome–year stratified "
            "paired bootstrap 95% CI and {}. NRI: categorical NRI; for two binary "
            "classifiers its components equal the differences in sensitivity and "
            "specificity."
        ).format(report["n_common"], report["n_excluded"], error_test)

    table_paths = {"docx": None, "csv": None}
    if export_tables:
        if figures_dir is None:
            raise ValueError("figures_dir is required when export_tables=True")
        tables_dir = os.path.join(figures_dir, "tables")
        os.makedirs(tables_dir, exist_ok=True)
        stem = os.path.join(tables_dir, "paired_comparison_model_vs_triage")
        table_paths = {"docx": stem + ".docx", "csv": stem + ".csv"}
        _export_word_table(display_df, title, footnote, table_paths["docx"])
        display_df.to_csv(table_paths["csv"], index=False)
        print("Paired comparison Word table saved: {}".format(table_paths["docx"]))
        print("Paired comparison CSV saved: {}".format(table_paths["csv"]))
        if paired_inference is not None:
            sex_path = os.path.join(tables_dir, "paired_comparison_sex_gaps.docx")
            matched_path = os.path.join(
                tables_dir, "paired_comparison_matched_operating_point.docx"
            )
            _export_word_table(
                paired_inference["sex_gaps"]["table"],
                "Sex-gap comparison between the ML model and telephone triage",
                "Gaps are Men minus Women on the test set. DD is the ML model gap "
                "minus the telephone triage gap. Because gaps can be negative or "
                "change sign, whether the disparity widens, narrows or reverses "
                "must be read from both gaps; DD alone does not indicate direction. "
                "DD 95% CI and p-values: outcome–year stratified paired bootstrap "
                "(1,000 resamples).",
                sex_path,
            )
            matched_result = paired_inference["matched_operating_point"]
            _export_word_table(
                matched_result["table"],
                "Model versus telephone triage at matched operating points",
                "Thresholds were selected exclusively from out-of-fold training "
                "probabilities and applied unchanged to the common test subset. At "
                "each threshold both sensitivity and specificity are compared; the "
                "matched metric is approximate on the test set because thresholds "
                "were selected on OOF training probabilities.",
                matched_path,
            )
            table_paths["sex_gaps_docx"] = sex_path
            table_paths["matched_operating_point_docx"] = matched_path
            print("Sex-gap Word table saved: {}".format(sex_path))
            print("Matched operating-point Word table saved: {}".format(matched_path))

    triage_matrix = {
        key.lower(): triage_metrics[key] for key in ("TN", "FP", "FN", "TP")
    }
    model_matrix = {
        key.lower(): model_metrics[key] for key in ("TN", "FP", "FN", "TP")
    }
    return {
        "report": report,
        "triage_confusion_matrix": triage_matrix,
        "model_confusion_matrix": model_matrix,
        "confusion_matrices": {"triage": triage_matrix, "model": model_matrix},
        "triage_metrics": triage_metrics,
        "model_metrics": model_metrics,
        "metrics": {"triage": triage_metrics, "model": model_metrics},
        "differences": differences,
        "operational_totals": operational_totals,
        "table": display_df,
        "table_paths": table_paths,
        "title": title,
        "footnote": footnote,
        "paired_inference": paired_inference,
    }
